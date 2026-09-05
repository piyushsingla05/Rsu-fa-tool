"""Regression suite. Run after every change.

The Computershare case is the anchor: the engine must keep reproducing a real
manual working paper. The others prove the supported-but-incomplete scenarios
still produce a workbook rather than an error.
"""
from __future__ import annotations

import datetime as dt
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = Path("/home/claude/brokers/RSU brokers/Computershare/"
              "PortfolioDetails_CLIENT_COMPUTERSHARE_01.xlsx")
TOLERANCE = 0.001      # 0.1%


def run(args: list[str]) -> tuple[bool, str]:
    p = subprocess.run([sys.executable, "-m", "src.run", *args],
                       cwd=ROOT, capture_output=True, text=True, timeout=600)
    return p.returncode == 0, (p.stdout + p.stderr)


def case(name: str, args: list[str], expect_tabs: list[str] | None = None):
    ok, out = run(args)
    status = "PASS" if ok else "FAIL"
    detail = ""
    if ok and expect_tabs:
        out_path = next((Path(a) for a in args if str(a).endswith(".xlsx")), None)
        if out_path and out_path.exists():
            names = load_workbook(out_path).sheetnames
            missing = [t for t in expect_tabs if t not in names]
            if missing:
                status, detail = "FAIL", f"missing tabs {missing}"
    if not ok:
        detail = out.strip().splitlines()[-1] if out.strip() else "no output"
    print(f"  [{status}] {name}{('  - ' + detail) if detail else ''}")
    return status == "PASS"


def computershare_anchor() -> bool:
    """The engine must still reproduce the client's own manual working file."""
    out = ROOT / "output" / "_regression_fy.xlsx"
    ok, msg = run(["clients/CLIENT_COMPUTERSHARE_01", "--from", "2025-04-01",
                   "--to", "2026-03-31", "--client", "_REG", "--out", str(out)])
    if not ok:
        print(f"  [FAIL] Computershare anchor - engine error: {msg.strip()[-160:]}")
        return False
    subprocess.run([sys.executable,
                    "/root/.claude/skills/synced/83b76784-fb84-4b5f-8a2b-86bbda13c6f3_"
                    "cfd19bfc-0475-47e5-aea4-3f53063f33d4/xlsx/scripts/recalc.py",
                    str(out), "240"], capture_output=True, timeout=600)

    raw = pd.read_excel(SOURCE, header=None)
    theirs = {pd.Timestamp(raw.iat[i, 0]).date(): (raw.iat[i, 17], raw.iat[i, 18],
                                                   raw.iat[i, 19])
              for i in range(7, 35) if pd.notna(raw.iat[i, 0])}
    ws = load_workbook(out, data_only=True)["FA_A3"]
    n = bad = 0
    worst = 0.0
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] is None:
            break
        d = r[6].date() if hasattr(r[6], "date") else r[6]
        if d not in theirs:
            continue
        n += 1
        for a, b in zip(theirs[d], (r[7], r[8], r[9])):
            e = abs(float(a) - float(b)) / max(abs(float(a)), 1)
            worst = max(worst, e)
            if e > TOLERANCE:
                bad += 1
    out.unlink(missing_ok=True)
    ok = n == 28 and bad == 0
    print(f"  [{'PASS' if ok else 'FAIL'}] Computershare anchor - {n} lots, "
          f"{bad} off by >0.1%, worst {worst:.4%}")
    return ok


TABS = ["SUMMARY", "FA_A2", "FA_A3", "RSU_MASTER", "VESTING_SALES",
        "DIVIDENDS_1042", "CAPITAL_GAINS", "FX_WORKING", "FSI_TR_WORKING",
        "RECONCILIATION", "REVIEW_REQUIRED"]

if __name__ == "__main__":
    print("RSU / FA engine regression")
    results = [
        computershare_anchor(),
        case("Computershare ESPP - calendar year",
             ["clients/CLIENT_COMPUTERSHARE_01", "--from", "2025-01-01",
              "--to", "2025-12-31", "--client", "_R1",
              "--out", "output/_r1.xlsx"], TABS),
        case("Limited statement + 1042-S only (Fidelity/MSFT)",
             ["clients/CLIENT_FIDELITY_01", "--from", "2025-01-01", "--to", "2025-12-31",
              "--client", "_R2", "--limited-statement", "--out", "output/_r2.xlsx"],
             TABS),
        case("Upload folder: unknown broker + portfolio Excel + user TTBR",
             ["clients/DEMO_UPLOAD", "--docs", "uploads/DEMO", "--from", "2025-01-01",
              "--to", "2025-12-31", "--client", "_R3", "--out", "output/_r3.xlsx"],
             TABS),
        case("Morgan Stanley - profiled PDF, 20:1 split, specific-lot cost",
             ["clients/CLIENT_MORGANSTANLEY_01", "--docs", "uploads/CLIENT_MORGANSTANLEY_01",
              "--from", "2025-01-01", "--to", "2025-12-31", "--client", "_R4",
              "--out", "output/_r4.xlsx"], TABS),
        case("Fidelity - documents only (SPS + 1042-S), limited statement",
             ["clients/CLIENT_FIDELITY_02", "--docs", "uploads/CLIENT_FIDELITY_02",
              "--from", "2025-01-01", "--to", "2025-12-31", "--client", "_R6",
              "--limited-statement", "--out", "output/_r6.xlsx"], TABS),
        case("Combined Fidelity + Schwab - transfer linkage",
             ["clients/CLIENT_COMBINED_01", "--docs", "uploads/CLIENT_COMBINED_01",
              "--from", "2025-01-01", "--to", "2025-12-31", "--client", "_R7",
              "--out", "output/_r7.xlsx"], TABS),
        case("E*TRADE - single quarterly statement, LIMITED DATA path",
             ["clients/CLIENT_ETRADE_01", "--docs", "uploads/CLIENT_ETRADE_01",
              "--from", "2025-01-01", "--to", "2025-12-31", "--client", "_R8",
              "--limited-statement", "--out", "output/_r8.xlsx"], TABS),
        case("Schwab - positions, cash, dividends+NRA tax, closed-lot gains",
             ["clients/CLIENT_SCHWAB_01", "--docs", "uploads/CLIENT_SCHWAB_01",
              "--from", "2025-01-01", "--to", "2025-12-31", "--client", "_R5",
              "--out", "output/_r5.xlsx"], TABS),
        case("E*TRADE G&L Expanded - adjusted capital gains, FY period",
             ["clients/ETRADE_GL", "--docs", "uploads/ETRADE_GL",
              "--from", "2026-04-01", "--to", "2027-03-31", "--client", "_R9",
              "--out", "output/_r9.xlsx"], TABS),
        case("E*TRADE G&L Expanded + Google Finance FX fallback",
             ["clients/ETRADE_GL", "--docs", "uploads/ETRADE_GL",
              "--from", "2026-04-01", "--to", "2027-03-31", "--client", "_R10",
              "--google-fx", "tests/fixtures/google_fx_intc.csv",
              "--out", "output/_r10.xlsx"], TABS),
        case("UBS - 5 statements per PDF, option exercise, dated 1042-S",
             ["clients/CLIENT_UBS_01", "--docs", "uploads/CLIENT_UBS_01",
              "--from", "2025-01-01", "--to", "2025-12-31", "--client", "_R11",
              "--out", "output/_r11.xlsx"], TABS),
    ]
    for extra, label in (("morgan_regression.py", "Morgan Stanley evidence regression"),
                         ("schwab_regression.py", "Schwab evidence regression"),
                         ("fidelity_regression.py",
                          "Fidelity + combined-client regression"),
                         ("etrade_regression.py",
                          "E*TRADE limited-data regression"),
                         ("etrade_gl_regression.py",
                          "E*TRADE G&L Expanded - adjusted capital gains"),
                         ("ubs_regression.py",
                          "UBS evidence regression"),
                         ("fx_fallback_regression.py",
                          "Universal FX fallback - SBI > Google Finance > blank"),
                         ("fx_thirdsource_regression.py",
                          "Third/fourth FX source - ECB and FBIL"),
                         ("fx_strict_mode_regression.py",
                          "Strict-mode fix - exact date only in the ranked pass"),
                         ("test_coverage_and_parity.py",
                          "Dividend coverage + generic parity"),
                         ("api_regression.py",
                          "UI / API - screen == engine == workbook")):
        r = subprocess.run([sys.executable, str(Path(__file__).parent / extra)],
                           cwd=ROOT, capture_output=True, text=True, timeout=900)
        results.append(r.returncode == 0)
        print(f"  [{'PASS' if r.returncode == 0 else 'FAIL'}] {label}")

    for f in ("_r1", "_r2", "_r3", "_r4", "_r5", "_r6", "_r7", "_r8", "_r9", "_r10", "_r11"):
        (ROOT / "output" / f"{f}.xlsx").unlink(missing_ok=True)
    print(f"\n{sum(results)}/{len(results)} passed")
    sys.exit(0 if all(results) else 1)
