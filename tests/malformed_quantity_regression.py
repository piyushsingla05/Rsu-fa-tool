"""Regression tests for the malformed-quantity hardening audit (HIGH finding).

Prior defect (investigated read-only, then confirmed): every quantity/shares
extraction site used G.to_number(v) with its default on_invalid="zero", then
gated the row with `if not qty: continue` (or equivalent). That idiom cannot
tell a genuinely blank/absent cell apart from one that was PRESENT but failed
to parse (garbled text, a stray footnote, "12.34.56") - both silently became
0.0 and the row vanished with no trace anywhere in the engine. Nine ingestion
sites had this shape:

  1. src/ingest/extract.py  _extract_pdf_sections role "holdings_as_acquisition"
  2. src/ingest/extract.py  _extract_pdf_sections role "option_exercises"
  3. src/ingest/extract.py  _extract_pdf_sections role "unvested_awards"
  4. src/ingest/extract.py  _extract_pdf_sections role "transfers"
  5. src/ingest/extract.py  _extract_pdf_sections role "sales"
  6. src/ingest/extract.py  _extract_pdf_sections role "closed_lot_gains"
  7. src/ingest/extract.py  _extract_closed_lot_table (spreadsheet twin of 6)
  8. src/ingest/extract.py  _extract_events (generic AND profiled fallback)
  9. src/parsers/computershare.py  ComputersharePlanParser._num()/parse()

Fix: each site now parses with on_invalid="nan", and an explicit `qty != qty`
check (never truthiness) diverts a malformed value to a NEW SourceRecord list,
rec.invalid_quantities, naming the role/event/symbol/broker/source/raw value,
instead of a bare `continue`. src/run.py's flag_invalid_quantities() (mirrors
the existing flag_closed_lot_reports()/recon_breaks pattern exactly) turns
every such entry into a register.blocker(INVALID_NUMERIC_FIELD, ...) - the
SAME diagnostic category compute.py's build_lots() already uses for this
class of defect, no new parallel framework invented. A genuinely blank/absent
value, and a genuine numeric zero, both still take the pre-existing silent
path unchanged.

A downstream gap was also closed: build_lots()'s DISPOSING_EVENTS branch
(SELL/TRANSFER_OUT) had no NaN-quantity guard at all (only ACQUIRING_EVENTS
did, from a prior commit) - a NaN reaching it would have silently corrupted
lot.remaining via NaN arithmetic, with NO diagnostic. It now mirrors the
ACQUIRING_EVENTS guard. The SPLIT ratio has the same new guard, and every
downstream aggregate quantity walk that used the unguarded `r["quantity"] or
0`/`or 1` idiom (which does NOT catch NaN, since NaN is truthy in Python) now
routes through a new _qty_or_zero() helper that excludes an already-flagged
NaN from its own sum instead of letting it silently propagate.

Entirely self-contained (synthetic tables/documents/events) - no private
data, no network, no change to any broker profile, the acceptance baseline,
FX logic, or the six prior commits.

Run: python3 -m tests.malformed_quantity_regression
"""
from __future__ import annotations

import datetime as dt
import math
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.ingest.extract import (                                       # noqa: E402
    SourceRecord, _extract_closed_lot_table, _extract_events, _extract_pdf_sections,
)
from src.ingest.profiles import Profile                                # noqa: E402
from src.ingest.tabular import Document, Table                         # noqa: E402
from src.parsers.computershare import ComputersharePlanParser          # noqa: E402
from src.compute import build_cg, build_lots, build_reconciliation, ComputeOptions  # noqa: E402
from src.fx import FXTable                                             # noqa: E402
from src.models import EVENTS_COLUMNS, Period                          # noqa: E402
from src.review import INVALID_NUMERIC_FIELD, Register                 # noqa: E402

PASS, FAIL = "PASS", "FAIL"
_fails = 0


def check(ok, label, detail=""):
    global _fails
    status = PASS if ok else FAIL
    print(f"  {status:<6} {label}{('  ' + detail) if detail else ''}")
    if not ok:
        _fails += 1


def section(title):
    print()
    print("=" * 92)
    print(title)
    print("=" * 92)


def is_nan(x):
    try:
        return isinstance(x, float) and math.isnan(x)
    except Exception:
        return False


def new_rec(period_end=dt.date(2024, 12, 31)):
    return SourceRecord(file="synthetic.pdf", kind="pdf", document_type="BROKER_STATEMENT",
                        profile_id="test", broker="TESTBROKER", confidence=1.0,
                        period_end=period_end)


def pdf_doc(lines):
    return Document(path=Path("synthetic.pdf"), kind="pdf",
                    text="\n".join(lines).lower(), lines=lines)


def pdf_profile(section_cfg, rules=None):
    return Profile(id="test", label="test", document_type="BROKER_STATEMENT",
                   layout={"kind": "pdf_sections", "sections": [section_cfg]},
                   rules=rules or {})


def invalid_qty_entries(rec, role):
    return [iv for iv in rec.invalid_quantities if iv["role"] == role]


def events_frame(rows):
    df = pd.DataFrame(rows)
    for c in EVENTS_COLUMNS:
        if c not in df.columns:
            df[c] = ""
    return df[EVENTS_COLUMNS]


def ev(date, broker, account, symbol, event, quantity, price_fc=0.0, **extra):
    row = {"date": date, "broker": broker, "account_no": account, "symbol": symbol,
           "event": event, "quantity": quantity, "price_fc": price_fc,
           "amount_fc": "", "tax_fc": "", "acquired_on": "", "cost_fc": "",
           "currency": "USD", "notes": ""}
    row.update(extra)
    return row


def coerced(rows):
    """Reproduce _read()'s CSV-events path: pd.to_numeric(errors="coerce")."""
    raw = pd.DataFrame(rows)
    raw["price_fc"] = pd.to_numeric(raw["price_fc"], errors="coerce")
    raw["quantity"] = pd.to_numeric(raw["quantity"], errors="coerce")
    return events_frame(raw.to_dict("records"))


def controlled_fx(register):
    fx = FXTable(ROOT / "config" / "fx_rates.csv", register=register)
    manual = pd.DataFrame([
        {"date": "2024-01-01", "currency": "USD", "ttbr": 83.0, "verified": "Y"},
        {"date": "2024-03-01", "currency": "USD", "ttbr": 83.2, "verified": "Y"},
        {"date": "2024-06-01", "currency": "USD", "ttbr": 83.8, "verified": "Y"},
    ])
    fx.merge(manual, priority=True)
    return fx


PERIOD = Period(dt.date(2024, 1, 1), dt.date(2024, 12, 31))


# ======================================================================
section("1. holdings_as_acquisition - malformed SHARES excluded, diagnosed")

profile1 = pdf_profile({
    "name": "positions", "role": "holdings_as_acquisition",
    "start_after": "POSITIONS",
    "row_regex": r"^POS (?P<ticker>[A-Z]+) (?P<shares>[\w.,$-]+) (?P<cost_total>[\d.,]+)$",
})
doc1 = pdf_doc(["POSITIONS", "POS AAPL 50 2000.00", "POS AAPL BADQTY 2000.00"])
rec1 = new_rec()
events1, unmapped1, holdings1 = _extract_pdf_sections(doc1, profile1, "TESTBROKER", rec1)

check(len(holdings1) == 1 and holdings1[0]["shares"] == 50.0,
      "only the valid-shares position becomes a holding - the malformed one "
      "is excluded, not fabricated as a zero-share position",
      f"{len(holdings1)} holding(s): {[h['shares'] for h in holdings1]}")
check(len(events1) == 1,
      "exactly one POSITION event row produced", f"{len(events1)} row(s)")
iv1 = invalid_qty_entries(rec1, "holdings_as_acquisition")
check(len(iv1) == 1 and iv1[0]["broker"] == "TESTBROKER" and iv1[0]["raw"] == "'BADQTY'",
      "a visible diagnostic names the broker, role and raw malformed value",
      f"{iv1}")


# ======================================================================
section("2. option_exercises - malformed QTY excluded, diagnosed")

profile2 = pdf_profile({
    "name": "exercises", "role": "option_exercises",
    "start_after": "EXERCISES",
    "row_regex": r"^EXER (?P<grant_no>\w+) (?P<exercise_date>\d{4}-\d{2}-\d{2}) (?P<qty>[\w.,-]+)$",
})
doc2 = pdf_doc(["EXERCISES", "EXER G1 2024-03-01 100", "EXER G2 2024-03-02 BADQTY"])
rec2 = new_rec()
events2, _, _ = _extract_pdf_sections(doc2, profile2, "TESTBROKER", rec2)

check(len(rec2.exercises) == 1 and rec2.exercises[0]["quantity"] == 100.0,
      "only the valid exercise is recorded - a dropped exercise would "
      "otherwise silently defeat the FMV-as-cost perquisite-double-tax fix",
      f"{len(rec2.exercises)} exercise(s)")
iv2 = invalid_qty_entries(rec2, "option_exercises")
check(len(iv2) == 1 and iv2[0]["event"] == "OPTION_EXERCISE",
      "a visible diagnostic is raised for the malformed exercise quantity",
      f"{iv2}")


# ======================================================================
section("3. unvested_awards - malformed quantity excluded, diagnosed")

profile3 = pdf_profile({
    "name": "unvested", "role": "unvested_awards",
    "start_after": "UNVESTED",
    "row_regex": r"^UNV (?P<award_no>\w+) (?P<unvested>[\w.,-]+)$",
})
doc3 = pdf_doc(["UNVESTED", "UNV A1 30", "UNV A2 BADQTY"])
rec3 = new_rec()
_extract_pdf_sections(doc3, profile3, "TESTBROKER", rec3)

check(len(rec3.unvested_awards) == 1 and rec3.unvested_awards[0]["unvested"] == 30.0,
      "only the valid unvested-award row is recorded", f"{rec3.unvested_awards}")
iv3 = invalid_qty_entries(rec3, "unvested_awards")
check(len(iv3) == 1, "a visible diagnostic is raised for the malformed award "
      "quantity", f"{iv3}")


# ======================================================================
section("4. transfers - malformed quantity excluded, diagnosed")

profile4 = pdf_profile({
    "name": "transfers", "role": "transfers", "direction": "OUT",
    "start_after": "TRANSFERS",
    "row_regex": r"^XFER (?P<date>\d{4}-\d{2}-\d{2}) (?P<qty>[\w.,-]+)$",
})
doc4 = pdf_doc(["TRANSFERS", "XFER 2024-04-01 40", "XFER 2024-04-02 BADQTY"])
rec4 = new_rec()
_extract_pdf_sections(doc4, profile4, "TESTBROKER", rec4)

check(len(rec4.transfers) == 1 and rec4.transfers[0]["quantity"] == 40.0,
      "only the valid transfer is recorded - a dropped malformed transfer "
      "would otherwise silently break build_cross_broker()'s FA-A2/A3 "
      "cross-broker quantity chain", f"{rec4.transfers}")
iv4 = invalid_qty_entries(rec4, "transfers")
check(len(iv4) == 1 and iv4[0]["event"] == "TRANSFER_OUT",
      "a visible diagnostic is raised, naming the transfer direction",
      f"{iv4}")


# ======================================================================
section("5. sales - malformed quantity excluded, diagnosed, no false SELL")

profile5 = pdf_profile({
    "name": "sales", "role": "sales",
    "start_after": "SALES",
    "row_regex": (r"^SALE (?P<date>\d{4}-\d{2}-\d{2}) (?P<qty>[\w.,-]+) "
                  r"(?P<price>[\d.,]+) (?P<amount>[\d.,]+)$"),
})
doc5 = pdf_doc(["SALES", "SALE 2024-05-01 20 50.00 1000.00",
               "SALE 2024-05-02 BADQTY 50.00 1000.00"])
rec5 = new_rec()
events5, _, _ = _extract_pdf_sections(doc5, profile5, "TESTBROKER", rec5)

check(len(events5) == 1 and events5.iloc[0]["quantity"] == 20.0,
      "only the valid SALE becomes a SELL event - no zero-quantity SELL is "
      "fabricated for the malformed row", f"{len(events5)} row(s)")
iv5 = invalid_qty_entries(rec5, "sales")
check(len(iv5) == 1 and iv5[0]["event"] == "SELL",
      "a visible diagnostic is raised for the malformed sale quantity",
      f"{iv5}")


# ======================================================================
section("6. closed_lot_gains (pdf_sections) - malformed quantity excluded")

profile6 = pdf_profile({
    "name": "closed_lots", "role": "closed_lot_gains",
    "start_after": "CLOSEDLOTS",
    "row_regex": (r"^CLG (?P<acquired>\d{4}-\d{2}-\d{2}) (?P<sold>\d{4}-\d{2}-\d{2}) "
                  r"(?P<qty>[\w.,-]+) (?P<proceeds>[\d.,]+) (?P<cost>[\d.,]+)$"),
})
doc6 = pdf_doc(["CLOSEDLOTS",
               "CLG 2024-01-01 2024-06-01 20 1400.00 1000.00",
               "CLG 2024-01-02 2024-06-02 BADQTY 1400.00 1000.00"])
rec6 = new_rec()
events6, _, _ = _extract_pdf_sections(doc6, profile6, "TESTBROKER", rec6)

check(len(events6) == 1 and events6.iloc[0]["quantity"] == 20.0,
      "only the valid closed-lot record becomes a SELL event",
      f"{len(events6)} row(s)")
iv6 = invalid_qty_entries(rec6, "closed_lot_gains")
check(len(iv6) == 1, "a visible diagnostic is raised for the malformed "
      "closed-lot quantity", f"{iv6}")


# ======================================================================
section("7. _extract_closed_lot_table (spreadsheet) - malformed quantity excluded")

profile7 = Profile(id="test7", label="test", document_type="BROKER_STATEMENT",
                   layout={}, rules={},
                   fields={"quantity": "Quantity", "sale_date": "Sale Date",
                           "proceeds_fc": "Proceeds", "cost_fc": "Cost",
                           "symbol": "Symbol"})
df7 = pd.DataFrame({
    "Quantity": [20, "BADQTY"],
    "Sale Date": ["2024-06-01", "2024-06-02"],
    "Proceeds": [1400.00, 1400.00],
    "Cost": [1000.00, 1000.00],
    "Symbol": ["AAPL", "AAPL"],
})
t7 = Table(df=df7, header_row=0, sheet="Sheet1", source_file="synthetic.xlsx", raw=None)
doc7 = Document(path=Path("synthetic.xlsx"), kind="excel", text="", tables=[t7])
rec7 = new_rec()
events7, _ = _extract_closed_lot_table(doc7, profile7, "TESTBROKER", rec7)

check(len(events7) == 1 and events7.iloc[0]["quantity"] == 20.0,
      "only the valid row becomes a SELL event - the malformed spreadsheet "
      "quantity is excluded, not zeroed", f"{len(events7)} row(s)")
check(len(rec7.closed_lots) == 1,
      "rec.closed_lots (the capital-gains-authority source) also excludes "
      "the malformed row", f"{len(rec7.closed_lots)} closed lot(s)")
iv7 = invalid_qty_entries(rec7, "closed_lot_gains_table")
check(len(iv7) == 1 and iv7[0]["symbol"] == "AAPL",
      "a visible diagnostic names the symbol for the malformed quantity",
      f"{iv7}")


# ======================================================================
section("8. _extract_events (generic + profiled fallback) - malformed "
        "quantity excluded for VEST/SELL/TRANSFER_OUT and DIVIDEND_REINVESTMENT")


def make_generic_doc(rows, extra_cols=None):
    headers = {"Date Acquired": [], "Quantity": [], "Price": [], "Symbol": [],
              "Description": []}
    for r in rows:
        headers["Date Acquired"].append(r.get("date", "01/01/2024"))
        headers["Quantity"].append(r.get("qty", ""))
        headers["Price"].append(r.get("price", "40.00"))
        headers["Symbol"].append(r.get("symbol", "AAPL"))
        headers["Description"].append(r.get("descriptor", ""))
    df = pd.DataFrame(headers)
    t = Table(df=df, header_row=0, sheet="Sheet1", source_file="synthetic.xlsx", raw=None)
    return Document(path=Path("synthetic.xlsx"), kind="excel", text="", tables=[t])


for label, descriptor, expect_event in [
    ("VEST (default classification)", "RSU Vesting", "VEST"),
    ("SELL (classified)", "Sale of shares", "SELL"),
    ("TRANSFER_OUT (classified)", "Transfer out to another broker", "TRANSFER_OUT"),
]:
    doc8 = make_generic_doc([{"qty": "BADQTY", "descriptor": descriptor}])
    rec8 = new_rec()
    events8, _ = _extract_events(doc8, None, "GENERIC_BROKER", rec8)
    check(len(events8) == 0,
          f"{label}: malformed quantity row excluded, not fabricated as a "
          "zero-quantity event", f"{len(events8)} row(s)")
    iv8 = invalid_qty_entries(rec8, "_extract_events")
    check(len(iv8) == 1 and iv8[0]["event"] == expect_event,
          f"{label}: a visible diagnostic names the classified event type",
          f"{iv8}")

# DIVIDEND_REINVESTMENT: a malformed quantity must not silently drop BOTH the
# BUY lot and its paired DIV income row (the Computershare double-omission
# risk, but here for ANY broker hitting the generic/profiled fallback).
doc8d = make_generic_doc([{"qty": "BADQTY", "descriptor": "Dividend Reinvestment Purchase"}])
rec8d = new_rec()
events8d, _ = _extract_events(doc8d, None, "GENERIC_BROKER", rec8d)
check(len(events8d) == 0,
      "DIVIDEND_REINVESTMENT: malformed quantity excludes BOTH the BUY lot "
      "and the paired DIV income row together - correct, since the true "
      "quantity is unknown - but with a diagnostic, never silently",
      f"{len(events8d)} row(s)")
iv8d = invalid_qty_entries(rec8d, "_extract_events")
check(len(iv8d) == 1, "a visible diagnostic is raised for the "
      "DIVIDEND_REINVESTMENT malformed quantity", f"{iv8d}")

# A genuinely valid quantity is completely unaffected by any of this.
doc8v = make_generic_doc([{"qty": "50", "descriptor": "RSU Vesting"}])
rec8v = new_rec()
events8v, _ = _extract_events(doc8v, None, "GENERIC_BROKER", rec8v)
check(len(events8v) == 1 and events8v.iloc[0]["quantity"] == 50.0,
      "a genuinely valid quantity still produces exactly one event, unchanged",
      f"{len(events8v)} row(s)")
check(len(rec8v.invalid_quantities) == 0,
      "no diagnostic is raised for a genuinely valid quantity", "")


# ======================================================================
section("9. Computershare parser - malformed quantity excludes BOTH the "
        "acquisition lot AND the paired dividend income, with a warning")

import openpyxl  # noqa: E402


def make_computershare_xlsx(path, qty_value, is_dividend=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Participant Name", "Jane Doe"])
    ws.append(["User ID", "ACC123"])
    ws.append([])
    ws.append(["Allocation date", "Plan", "Contribution Type", "Instrument",
              "Strike Price / Cost Basis", "Outstanding Quantity"])
    instr = "Dividend Shares" if is_dividend else "Purchase Shares"
    ws.append(["2024-03-01", "Microsoft ESPP", "ESPP", instr, 40.0, qty_value])
    wb.save(path)


cs_dir = Path("/tmp")

cs_valid = cs_dir / "malformed_qty_cs_valid.xlsx"
make_computershare_xlsx(cs_valid, 25.0)
data_valid = ComputersharePlanParser().parse(cs_valid)
check(len(data_valid.events) == 1 and data_valid.events.iloc[0]["quantity"] == 25.0,
      "a genuinely valid Computershare quantity is unaffected", "")

cs_bad = cs_dir / "malformed_qty_cs_bad_purchase.xlsx"
make_computershare_xlsx(cs_bad, "BADQTY", is_dividend=False)
data_bad = ComputersharePlanParser().parse(cs_bad)
check(len(data_bad.events) == 0,
      "Purchase Shares: malformed quantity excludes the acquisition lot, "
      "not fabricated as a zero-share BUY", f"{len(data_bad.events)} event(s)")
check(any("unparseable" in w for w in data_bad.warnings),
      "a visible warning is raised for the malformed Purchase Shares quantity",
      f"{data_bad.warnings}")

cs_bad_div = cs_dir / "malformed_qty_cs_bad_dividend.xlsx"
make_computershare_xlsx(cs_bad_div, "BADQTY", is_dividend=True)
data_bad_div = ComputersharePlanParser().parse(cs_bad_div)
check(len(data_bad_div.events) == 0,
      "Dividend Shares: malformed quantity excludes BOTH the acquisition lot "
      "AND the paired dividend-income row together, not just one of them",
      f"{len(data_bad_div.events)} event(s)")
check(any("unparseable" in w for w in data_bad_div.warnings),
      "a visible warning is raised for the malformed Dividend Shares quantity",
      f"{data_bad_div.warnings}")


# ======================================================================
section("10. build_lots() DISPOSING_EVENTS - NaN quantity never poisons "
        "lot.remaining, no false gain, INVALID_NUMERIC_FIELD raised")

events10 = coerced([
    ev(dt.date(2024, 1, 1), "X", "1", "AAPL", "VEST", "100", "40.0"),
    ev(dt.date(2024, 3, 1), "X", "1", "AAPL", "SELL", "garbled_qty", "50.0"),
    ev(dt.date(2024, 6, 1), "X", "1", "AAPL", "SELL", "30", "60.0"),
])
reg10 = Register()
lots10, matches10 = build_lots(events10, reg10, as_at=dt.date(2024, 12, 31))

check(len(lots10) == 1, "the malformed SELL builds no SaleMatch of its own "
      "and consumes no lot - only the legitimate 30-share sale does",
      f"{len(lots10)} lot(s), {len(matches10)} match(es)")
check(not is_nan(lots10[0].remaining) and abs(lots10[0].remaining - 70.0) < 1e-9,
      "lot.remaining is a clean 70.0 (100 vested - 30 legitimately sold) - "
      "NEVER NaN, even though a malformed disposal was seen for this symbol",
      f"remaining={lots10[0].remaining!r}")
check(len(matches10) == 1 and abs(matches10[0].quantity - 30.0) < 1e-9,
      "exactly one real SaleMatch exists, for the legitimate 30-share sale "
      "only - the malformed row created no phantom match/gain",
      f"{[(m.quantity, m.sale_price_fc) for m in matches10]}")

blockers10 = [f for f in reg10.blockers if f.reason == INVALID_NUMERIC_FIELD]
check(len(blockers10) == 1 and "AAPL" in blockers10[0].subject,
      "an explicit INVALID_NUMERIC_FIELD blocker is raised for the "
      "unparseable disposal quantity", blockers10[0].detail if blockers10 else "none")

fx10 = controlled_fx(reg10)
opts10 = ComputeOptions()
cg10 = build_cg(matches10, fx10, opts10, PERIOD, reg10)
check(len(cg10) == 1,
      "capital gains has exactly one row - no false gain fabricated from the "
      "malformed disposal", f"{len(cg10)} row(s)")


# ======================================================================
section("11. build_lots() TRANSFER_OUT (also DISPOSING_EVENTS) - same guard")

events11 = coerced([
    ev(dt.date(2024, 1, 1), "X", "1", "MSFT", "VEST", "100", "40.0"),
    ev(dt.date(2024, 3, 1), "X", "1", "MSFT", "TRANSFER_OUT", "garbled_qty", "0.0"),
])
reg11 = Register()
lots11, matches11 = build_lots(events11, reg11, as_at=dt.date(2024, 12, 31))
check(not is_nan(lots11[0].remaining) and abs(lots11[0].remaining - 100.0) < 1e-9,
      "a malformed TRANSFER_OUT quantity leaves the lot at its full, "
      "un-poisoned 100.0 remaining rather than NaN", f"{lots11[0].remaining!r}")
blockers11 = [f for f in reg11.blockers if f.reason == INVALID_NUMERIC_FIELD]
check(len(blockers11) == 1, "an explicit INVALID_NUMERIC_FIELD blocker is "
      "raised for the malformed TRANSFER_OUT quantity", "")


# ======================================================================
section("12. build_lots() SPLIT - malformed ratio never poisons open lots")

events12 = coerced([
    ev(dt.date(2024, 1, 1), "X", "1", "GOOG", "VEST", "10", "100.0"),
    ev(dt.date(2024, 2, 1), "X", "1", "GOOG", "SPLIT", "garbled_ratio"),
])
reg12 = Register()
lots12, _ = build_lots(events12, reg12, as_at=dt.date(2024, 12, 31))
check(not is_nan(lots12[0].quantity) and not is_nan(lots12[0].remaining)
      and not is_nan(lots12[0].price_fc),
      "a malformed SPLIT ratio leaves the existing lot's quantity/remaining/"
      "price completely un-poisoned (no NaN multiplication applied)",
      f"quantity={lots12[0].quantity!r} remaining={lots12[0].remaining!r} "
      f"price_fc={lots12[0].price_fc!r}")
blockers12 = [f for f in reg12.blockers if f.reason == INVALID_NUMERIC_FIELD]
check(len(blockers12) == 1, "an explicit INVALID_NUMERIC_FIELD blocker is "
      "raised for the malformed SPLIT ratio", "")


# ======================================================================
section("13. Reconciliation/aggregate quantity walks - NaN never propagates")

events13 = coerced([
    ev(dt.date(2024, 1, 1), "X", "1", "AAPL", "VEST", "100", "40.0"),
    ev(dt.date(2024, 3, 1), "X", "1", "AAPL", "SELL", "garbled_qty", "50.0"),
    ev(dt.date(2024, 6, 1), "X", "1", "AAPL", "SELL", "30", "60.0"),
])
reg13 = Register()
lots13, _ = build_lots(events13, reg13, as_at=dt.date(2024, 12, 31))
recon13 = build_reconciliation(events13, lots13, PERIOD, reg13)
row13 = recon13[recon13["Symbol"] == "AAPL"].iloc[0]
check(all(not is_nan(row13[c]) for c in
          ("Opening Quantity", "Sold", "Expected Closing", "Closing per Lots",
           "Difference")),
      "every reconciliation figure is a real number - the malformed SELL's "
      "quantity never propagates NaN into the reconciliation walk",
      f"{dict(row13)}")
check(row13["Status"] == "Reconciled",
      "the walk stays reconciled - build_lots() and build_reconciliation() "
      "now consistently exclude the SAME malformed row (both via the "
      "explicit qty != qty guard), so they still agree with each other; the "
      "malformed value is not silently lost, it is separately flagged by "
      "build_lots()'s own INVALID_NUMERIC_FIELD blocker (section 10 above)",
      f"status={row13['Status']!r}")


# ======================================================================
print()
print("=" * 92)
if _fails:
    print(f"RESULT: {_fails} CHECK(S) FAILED")
else:
    print("RESULT: ALL CHECKS PASSED")
print("=" * 92)

if __name__ == "__main__":
    sys.exit(1 if _fails else 0)
